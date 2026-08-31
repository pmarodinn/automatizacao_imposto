"""Estilos do Excel usados pelo gerador de planilha.

Concentra fontes, bordas e alinhamentos em um só lugar para que a aparência da
planilha continue igual à do modelo original e possa ser ajustada por
`config/layout.yaml` sem espalhar formatação pelo código.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FINA = Side(style="thin")
MEDIA = Side(style="medium")

BORDA_COMPLETA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
BORDA_SUPERIOR = Border(top=FINA)
SEM_BORDA = Border()


@dataclass
class Tema:
    """Fonte base da planilha e tamanhos derivados."""

    fonte: str = "Arial"
    tamanho: float = 10.0

    def font(
        self,
        *,
        nome: str | None = None,
        tamanho: float | None = None,
        negrito: bool = False,
        italico: bool = False,
        sublinhado: str | None = None,
    ) -> Font:
        return Font(
            name=nome or self.fonte,
            size=tamanho or self.tamanho,
            bold=negrito,
            italic=italico,
            underline=sublinhado,
        )

    @classmethod
    def de_layout(cls, estilo: dict[str, Any]) -> "Tema":
        return cls(
            fonte=estilo.get("fonte", "Arial"),
            tamanho=float(estilo.get("fonte_tamanho", 10)),
        )


def alinhamento(horizontal: str = "center", *, quebra: bool = False, vertical: str = "center") -> Alignment:
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=quebra)


def aplicar_borda(ws: Worksheet, primeira_linha: int, ultima_linha: int, colunas: int) -> None:
    """Desenha a grade da tabela em todo o retângulo indicado."""
    for linha in range(primeira_linha, ultima_linha + 1):
        for coluna in range(1, colunas + 1):
            ws.cell(row=linha, column=coluna).border = BORDA_COMPLETA


def mesclar(ws: Worksheet, linha: int, coluna: int, span: int) -> None:
    if span > 1:
        ws.merge_cells(
            start_row=linha, start_column=coluna, end_row=linha, end_column=coluna + span - 1
        )


def definir_larguras(ws: Worksheet, larguras: dict[int, float]) -> None:
    for indice, largura in larguras.items():
        ws.column_dimensions[get_column_letter(indice)].width = largura


def altura_estimada(texto: str, caracteres_por_linha: int, minima: float, maxima: float) -> float:
    """Altura da linha proporcional ao tamanho do texto que precisa caber."""
    if not texto:
        return minima
    linhas = max(1, -(-len(texto) // max(1, caracteres_por_linha)))
    return min(maxima, max(minima, 12.0 + 11.5 * linhas))
