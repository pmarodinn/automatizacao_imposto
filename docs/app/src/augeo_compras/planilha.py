"""Geração do arquivo Excel: aba de controle + uma aba por proforma.

Reproduz o modelo "TOTAL PRICE AND PROFORMA", mas com o número de linhas do
pedido e com as fórmulas ligadas entre as abas, do mesmo jeito que o original:

    aba de controle (NET PRICE)      aba(s) PROFORMA
    ---------------------------      -----------------------------------------
    quantidade e preço reais    -->  puxa da aba de controle (ou valor próprio,
    TOTAL / FREIGHT                  quando a proforma declara algo diferente)
    Invoice Value  <-------------->  SUBTOTAL + FREIGHT = TOTAL
    Service Value = o que sobra
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import Coluna, Config, Fornecedor
from .estilos import (
    BORDA_COMPLETA,
    BORDA_SUPERIOR,
    Tema,
    alinhamento,
    altura_estimada,
    definir_larguras,
    mesclar,
)
from .pedido import PedidoResolvido, ProformaResolvida, data_por_extenso

LINHAS_LOGO = 6           # linhas reservadas às logos, no topo da proforma
LINHAS_IDENTIFICACAO = 4  # PROFORMA: / número / data / responsável
ALTURA_LINHA_TIMBRE = 15.75  # pontos
PIXELS_POR_PONTO = 96 / 72
PAPEL = {"A4": 9, "LETTER": 1, "A3": 8}


class GeradorPlanilha:
    """Monta o workbook a partir de um pedido já resolvido."""

    def __init__(self, config: Config):
        self.config = config
        self.layout = config.layout
        self.comercial = config.comercial
        self.empresa = config.empresa
        self.tema = Tema.de_layout(self.layout.estilo)

        # Posição inicial (1-based) de cada coluna lógica na planilha.
        self.posicoes: dict[str, int] = {}
        coluna_atual = 1
        for coluna in self.layout.colunas:
            self.posicoes[coluna.chave] = coluna_atual
            coluna_atual += coluna.span
        self.total_colunas = coluna_atual - 1
        self.coluna_auxiliar = coluna_atual  # primeira coluna livre à direita

    # ------------------------------------------------------------------ #
    # Utilidades de layout
    # ------------------------------------------------------------------ #
    def letra(self, chave: str) -> str:
        return get_column_letter(self.posicoes[chave])

    def _rotulo(self, coluna: Coluna) -> str:
        return coluna.rotulo.format(
            moeda=self.comercial.moeda, simbolo=self.comercial.simbolo_moeda
        )

    def _aplicar_larguras(self, ws: Worksheet) -> None:
        larguras: dict[int, float] = {}
        for coluna in self.layout.colunas:
            base = self.posicoes[coluna.chave]
            for deslocamento in range(coluna.span):
                if deslocamento < len(coluna.larguras):
                    larguras[base + deslocamento] = coluna.larguras[deslocamento]
        definir_larguras(ws, larguras)

    def _configurar_pagina(self, ws: Worksheet) -> None:
        pagina = self.layout.pagina
        ws.page_setup.orientation = pagina.get("orientacao", "portrait")
        ws.page_setup.paperSize = PAPEL.get(str(pagina.get("papel", "A4")).upper(), 9)
        if pagina.get("ajustar_largura", True):
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
        if pagina.get("escala"):
            ws.page_setup.scale = int(pagina["escala"])
        ws.print_options.horizontalCentered = True

    def _cabecalho_tabela(self, ws: Worksheet, linha: int) -> None:
        estilo = self.layout.estilo
        ws.row_dimensions[linha].height = float(estilo.get("altura_linha_cabecalho", 25.4))
        for coluna in self.layout.colunas:
            indice = self.posicoes[coluna.chave]
            mesclar(ws, linha, indice, coluna.span)
            celula = ws.cell(row=linha, column=indice, value=self._rotulo(coluna))
            celula.font = self.tema.font(
                nome=estilo.get("fonte_cabecalho_tabela"),
                tamanho=float(estilo.get("tamanho_cabecalho_tabela", 10)),
                negrito=True,
                italico=True,
            )
            celula.alignment = alinhamento("center", quebra=True, vertical="center")
        for indice in range(1, self.total_colunas + 1):
            ws.cell(row=linha, column=indice).border = BORDA_COMPLETA

    def _escrever_item(
        self,
        ws: Worksheet,
        linha: int,
        valores: dict[str, Any],
        descricao_para_altura: str,
    ) -> None:
        """Escreve uma linha de item aplicando o estilo de cada coluna."""
        estilo = self.layout.estilo
        ws.row_dimensions[linha].height = altura_estimada(
            descricao_para_altura,
            int(estilo.get("caracteres_por_linha_descricao", 26)),
            float(estilo.get("altura_linha_item_minima", 30)),
            float(estilo.get("altura_linha_item_maxima", 75)),
        )
        for coluna in self.layout.colunas:
            indice = self.posicoes[coluna.chave]
            mesclar(ws, linha, indice, coluna.span)
            celula = ws.cell(row=linha, column=indice, value=valores.get(coluna.chave))
            celula.font = self.tema.font(
                nome=coluna.fonte,
                tamanho=coluna.fonte_tamanho,
                negrito=coluna.negrito,
            )
            celula.alignment = alinhamento(
                coluna.alinhamento, quebra=coluna.quebra_linha, vertical="center"
            )
            celula.number_format = coluna.formato
        for indice in range(1, self.total_colunas + 1):
            ws.cell(row=linha, column=indice).border = BORDA_COMPLETA

    def _bloco_total(
        self,
        ws: Worksheet,
        linha: int,
        rotulo: str,
        valor: Any,
        *,
        coluna_rotulo: str = "codigo",
        coluna_valor: str = "preco",
        formato: str = "#,##0.00",
    ) -> None:
        """Escreve uma linha de totalização (rótulo + valor, com borda)."""
        indice_rotulo = self.posicoes[coluna_rotulo]
        celula_rotulo = ws.cell(row=linha, column=indice_rotulo, value=rotulo)
        celula_rotulo.font = self.tema.font(negrito=True)
        celula_rotulo.alignment = alinhamento("center")
        celula_rotulo.border = BORDA_COMPLETA

        indice_valor = self.posicoes[coluna_valor]
        celula_valor = ws.cell(row=linha, column=indice_valor, value=valor)
        celula_valor.font = self.tema.font(negrito=True)
        celula_valor.alignment = alinhamento("right")
        celula_valor.number_format = formato
        celula_valor.border = BORDA_COMPLETA

    # ------------------------------------------------------------------ #
    # Aba de controle
    # ------------------------------------------------------------------ #
    def _aba_controle(self, wb: Workbook, resolvido: PedidoResolvido) -> dict[str, Any]:
        ws = wb.create_sheet(self.layout.abas.get("controle", "NET PRICE"))
        self._aplicar_larguras(ws)
        self._configurar_pagina(ws)

        pedido = resolvido.pedido
        titulo = ws.cell(row=1, column=1, value="CONTROL")
        titulo.font = self.tema.font(tamanho=12, negrito=True)
        titulo.alignment = alinhamento("center")
        titulo.border = BORDA_COMPLETA

        identificacao = " | ".join(
            parte
            for parte in (
                pedido.referencia,
                pedido.cliente,
                pedido.data.strftime("%d/%m/%Y"),
            )
            if parte
        )
        if identificacao:
            celula = ws.cell(row=1, column=self.posicoes["part_number"], value=identificacao)
            celula.font = self.tema.font(negrito=True)
            celula.alignment = alinhamento("left")

        linha_cabecalho = 3
        self._cabecalho_tabela(ws, linha_cabecalho)

        primeira = linha_cabecalho + 1
        col_qtd, col_preco, col_total = (self.letra(c) for c in ("quantidade", "preco", "total"))

        for deslocamento, item in enumerate(resolvido.linhas):
            linha = primeira + deslocamento
            self._escrever_item(
                ws,
                linha,
                {
                    "posicao": item.posicao,
                    "quantidade": item.quantidade,
                    "part_number": item.part_number,
                    "descricao": item.descricao,
                    "codigo": item.tipo,
                    "preco": item.preco_unitario,
                    "total": f"={col_preco}{linha}*{col_qtd}{linha}",
                },
                item.descricao,
            )
            if self.comercial.coluna_auxiliar.ativa:
                aux = ws.cell(
                    row=linha,
                    column=self.coluna_auxiliar,
                    value=f"={col_preco}{linha}*{self.comercial.coluna_auxiliar.fator}",
                )
                aux.number_format = "0.00"

        # Linhas em branco reservadas para acréscimos manuais.
        reservadas = max(0, self.comercial.linhas_reservadas - len(resolvido.linhas))
        posicao = (
            resolvido.linhas[-1].posicao + self.comercial.posicao_passo
            if resolvido.linhas
            else self.comercial.posicao_inicial
        )
        for deslocamento in range(reservadas):
            linha = primeira + len(resolvido.linhas) + deslocamento
            # As linhas em branco já vêm com a fórmula do total, para que basta
            # digitar quantidade e preço direto no Excel.
            self._escrever_item(
                ws,
                linha,
                {"posicao": posicao, "total": f"={col_preco}{linha}*{col_qtd}{linha}"},
                "",
            )
            if self.comercial.coluna_auxiliar.ativa:
                aux = ws.cell(
                    row=linha,
                    column=self.coluna_auxiliar,
                    value=f"={col_preco}{linha}*{self.comercial.coluna_auxiliar.fator}",
                )
                aux.number_format = "0.00"
            posicao += self.comercial.posicao_passo

        ultima = primeira + len(resolvido.linhas) + reservadas - 1
        if ultima < primeira:  # pedido vazio
            ultima = primeira - 1

        if self.comercial.coluna_auxiliar.ativa:
            ws.column_dimensions[get_column_letter(self.coluna_auxiliar)].hidden = True

        linha_total = ultima + 1
        rotulos = self.layout.totais
        self._bloco_total(
            ws,
            linha_total,
            rotulos.get("total_controle", "TOTAL"),
            f"=SUM({col_total}{primeira}:{col_total}{ultima})" if ultima >= primeira else 0,
            coluna_valor="total",
        )

        linha_frete = linha_total + 1
        self._bloco_total(ws, linha_frete, rotulos.get("frete", "FREIGHT"), None)

        linha_faturado = linha_frete + 2
        linha_servico = linha_faturado + 1
        linha_geral = linha_servico + 1
        self._bloco_total(ws, linha_faturado, rotulos.get("valor_faturado", "Invoice Value"), None)
        self._bloco_total(ws, linha_servico, rotulos.get("valor_servico", "Service Value"), None)
        self._bloco_total(ws, linha_geral, rotulos.get("total_geral", "Total"), None)

        for linha in (linha_total, linha_frete, linha_faturado, linha_servico, linha_geral):
            ws.row_dimensions[linha].height = 15.5
        for coluna in range(1, self.posicoes["codigo"]):
            ws.cell(row=linha_total, column=coluna).border = BORDA_SUPERIOR

        return {
            "aba": ws,
            "primeira": primeira,
            "ultima": ultima,
            "linha_total": linha_total,
            "linha_frete": linha_frete,
            "linha_faturado": linha_faturado,
            "linha_servico": linha_servico,
            "linha_geral": linha_geral,
        }

    # ------------------------------------------------------------------ #
    # Aba de proforma
    # ------------------------------------------------------------------ #
    def _inserir_logo(
        self, ws: Worksheet, caminho: Path, ancora: str, largura_max: int
    ) -> None:
        """Insere a logo ajustada à largura pedida e à altura reservada.

        A altura manda: sem isso, uma imagem alta invadiria as linhas de texto
        do timbre e cobriria o endereço ou o número da proforma.
        """
        try:
            from openpyxl.drawing.image import Image

            logo = Image(str(caminho))
            if not logo.width or not logo.height:
                return
            altura_max = LINHAS_LOGO * ALTURA_LINHA_TIMBRE * PIXELS_POR_PONTO
            escala = min(largura_max / logo.width, altura_max / logo.height)
            logo.width = max(1, int(logo.width * escala))
            logo.height = max(1, int(logo.height * escala))
            ws.add_image(logo, ancora)
        except Exception:  # Pillow ausente ou imagem inválida: segue sem logo
            pass

    def _timbre(
        self,
        ws: Worksheet,
        proforma: ProformaResolvida,
        fornecedor: Fornecedor | None,
    ) -> int:
        """Bloco superior da proforma. Devolve a última linha usada.

        À esquerda, a logo e os dados de quem emite; à direita, a logo do
        fornecedor de quem se está comprando e a identificação da proforma.
        """
        coluna_direita = max(2, self.posicoes["descricao"])
        textos_empresa = self.empresa.linhas_timbre()
        primeira_linha_texto = LINHAS_LOGO + 1
        ultima_linha = max(
            primeira_linha_texto + len(textos_empresa) - 1,
            primeira_linha_texto + LINHAS_IDENTIFICACAO - 1,
        )
        for linha in range(1, ultima_linha + 1):
            ws.row_dimensions[linha].height = 15.75

        # Logo de quem emite a proforma (à esquerda) e do fornecedor (à direita).
        if self.empresa.caminho_logo:
            self._inserir_logo(ws, self.empresa.caminho_logo, "A1", self.empresa.largura_logo)
        if fornecedor and fornecedor.caminho_logo:
            self._inserir_logo(
                ws,
                fornecedor.caminho_logo,
                f"{get_column_letter(coluna_direita)}1",
                fornecedor.largura_logo,
            )

        for deslocamento, texto in enumerate(textos_empresa):
            linha = primeira_linha_texto + deslocamento
            mesclar(ws, linha, 1, coluna_direita - 1)
            celula = ws.cell(row=linha, column=1, value=texto)
            celula.font = self.tema.font(tamanho=11, negrito=(deslocamento == 0))
            celula.alignment = alinhamento("left")

        blocos = [
            ("PROFORMA:", True),
            (proforma.numero, False),
            (data_por_extenso(proforma.data, self.empresa.cidade), False),
            (self.empresa.responsavel, False),
        ]
        for deslocamento, (texto, negrito) in enumerate(blocos):
            if not texto:
                continue
            linha = primeira_linha_texto + deslocamento
            mesclar(ws, linha, coluna_direita, self.total_colunas - coluna_direita + 1)
            celula = ws.cell(row=linha, column=coluna_direita, value=texto)
            celula.font = self.tema.font(tamanho=11, negrito=negrito)
            celula.alignment = alinhamento("left")

        return ultima_linha

    def _condicoes(self, ws: Worksheet, linha: int, proforma: ProformaResolvida) -> int:
        """Bloco "Condições gerais" no rodapé da proforma."""
        rodape = self.layout.rodape
        condicoes = proforma.condicoes
        coluna_valor = min(int(rodape.get("coluna_valor", 5) or 5), self.total_colunas)

        titulo = ws.cell(row=linha, column=1, value=rodape.get("titulo", "Condições gerais"))
        titulo.font = self.tema.font(tamanho=11, negrito=True)
        linha += 1

        def par(rotulo: str, valores: list[str], *, sublinhar: bool = True) -> None:
            nonlocal linha
            valores = [v for v in valores if v]
            if not valores:
                return
            celula = ws.cell(row=linha, column=1, value=rotulo)
            celula.font = self.tema.font(tamanho=11, sublinhado="single" if sublinhar else None)
            for deslocamento, valor in enumerate(valores):
                destino = ws.cell(row=linha + deslocamento, column=coluna_valor, value=valor)
                destino.font = self.tema.font(tamanho=11, negrito=(deslocamento == 0))
                destino.alignment = alinhamento("left")
            linha += len(valores)

        par(rodape.get("delivery_terms", "DELIVERY TERMS"), [condicoes.delivery_terms], sublinhar=False)
        par(rodape.get("payment_terms", "PAYMENT TERMS:"), [condicoes.payment_terms, *condicoes.banco])
        linha += 1
        par(rodape.get("shipment_terms", "SHIPMENT TERMS:"), [condicoes.shipment_terms, condicoes.modal])
        linha += 1
        if condicoes.agent_export:
            celula = ws.cell(row=linha, column=1, value=rodape.get("agent_export", "AGENT EXPORT"))
            celula.font = self.tema.font(tamanho=11, sublinhado="single")
            linha += 1
            destino = ws.cell(row=linha, column=1, value=condicoes.agent_export)
            destino.font = self.tema.font(negrito=True)
            linha += 1
        return linha

    def _aba_proforma(
        self,
        wb: Workbook,
        resolvido: PedidoResolvido,
        proforma: ProformaResolvida,
        controle: dict[str, Any],
        total_proformas: int,
        fornecedor: Fornecedor | None,
    ) -> dict[str, Any]:
        nome_base = self.layout.abas.get("proforma", "PROFORMA")
        nome = proforma.titulo_aba or (
            nome_base if total_proformas == 1 else f"{nome_base} {proforma.indice}"
        )
        ws = wb.create_sheet(nome[:31])
        self._aplicar_larguras(ws)
        self._configurar_pagina(ws)
        fim_timbre = self._timbre(ws, proforma, fornecedor)

        linha_cabecalho = fim_timbre + 2
        self._cabecalho_tabela(ws, linha_cabecalho)

        aba_controle = controle["aba"].title
        referencia_controle = f"'{aba_controle}'!"
        col_qtd, col_preco, col_total = (self.letra(c) for c in ("quantidade", "preco", "total"))
        primeira = linha_cabecalho + 1

        linha_no_controle = {
            id(item): controle["primeira"] + posicao
            for posicao, item in enumerate(resolvido.linhas)
        }
        for deslocamento, item in enumerate(proforma.linhas):
            linha = primeira + deslocamento
            linha_controle = linha_no_controle[id(item)]
            # Quando a proforma declara o mesmo que o controle, mantém a ligação
            # entre as abas; quando declara algo diferente, grava o valor próprio.
            quantidade = (
                f"={referencia_controle}{col_qtd}{linha_controle}"
                if abs(item.quantidade - item.quantidade_proforma) < 1e-9
                else item.quantidade_proforma
            )
            preco = (
                f"={referencia_controle}{col_preco}{linha_controle}"
                if abs(item.preco_unitario - item.preco_proforma) < 1e-9
                else item.preco_proforma
            )
            self._escrever_item(
                ws,
                linha,
                {
                    "posicao": item.posicao,
                    "quantidade": quantidade,
                    "part_number": f"={referencia_controle}{self.letra('part_number')}{linha_controle}",
                    "descricao": f"={referencia_controle}{self.letra('descricao')}{linha_controle}",
                    "codigo": f"={referencia_controle}{self.letra('codigo')}{linha_controle}",
                    "preco": preco,
                    "total": f"={col_preco}{linha}*{col_qtd}{linha}",
                },
                item.descricao,
            )

        ultima = primeira + len(proforma.linhas) - 1
        rotulos = self.layout.totais

        linha_subtotal = max(ultima, primeira - 1) + 1
        self._bloco_total(
            ws,
            linha_subtotal,
            rotulos.get("subtotal", "SUBTOTAL"),
            f"=SUM({col_total}{primeira}:{col_total}{ultima})" if ultima >= primeira else 0,
        )
        linha_frete = linha_subtotal + 1
        self._bloco_total(
            ws, linha_frete, rotulos.get("frete", "FREIGHT"), proforma.frete or None
        )
        linha_total = linha_frete + 1
        col_valor = self.letra("preco")
        self._bloco_total(
            ws,
            linha_total,
            rotulos.get("total", "TOTAL"),
            f"={col_valor}{linha_subtotal}+{col_valor}{linha_frete}",
        )
        for linha in (linha_subtotal, linha_frete, linha_total):
            ws.row_dimensions[linha].height = 15.5
        for coluna in range(1, self.posicoes["codigo"]):
            ws.cell(row=linha_subtotal, column=coluna).border = BORDA_SUPERIOR

        self._condicoes(ws, linha_total + 2, proforma)
        ws.print_area = f"A1:{get_column_letter(self.total_colunas)}{linha_total + 14}"

        return {"aba": ws, "linha_frete": linha_frete, "linha_total": linha_total}

    # ------------------------------------------------------------------ #
    # Conciliação entre as abas
    # ------------------------------------------------------------------ #
    def _conciliar(
        self, controle: dict[str, Any], proformas: list[dict[str, Any]]
    ) -> None:
        """Preenche FREIGHT / Invoice / Service da aba de controle."""
        ws = controle["aba"]
        col_valor = self.letra("preco")
        col_total = self.letra("total")

        def referencia(info: dict[str, Any], linha: int) -> str:
            return f"'{info['aba'].title}'!{col_valor}{linha}"

        fretes = "+".join(referencia(p, p["linha_frete"]) for p in proformas)
        totais = "+".join(referencia(p, p["linha_total"]) for p in proformas)

        ws.cell(row=controle["linha_frete"], column=self.posicoes["preco"]).value = (
            f"={fretes}" if fretes else 0
        )
        ws.cell(row=controle["linha_faturado"], column=self.posicoes["preco"]).value = (
            f"={totais}" if totais else 0
        )
        ws.cell(row=controle["linha_servico"], column=self.posicoes["preco"]).value = (
            f"={col_total}{controle['linha_total']}"
            f"-{col_valor}{controle['linha_faturado']}"
            f"+{col_valor}{controle['linha_frete']}"
        )
        ws.cell(row=controle["linha_geral"], column=self.posicoes["preco"]).value = (
            f"={col_valor}{controle['linha_faturado']}+{col_valor}{controle['linha_servico']}"
        )

    # ------------------------------------------------------------------ #
    # API
    # ------------------------------------------------------------------ #
    def fornecedor_do_pedido(self, resolvido: PedidoResolvido) -> Fornecedor | None:
        """Decide qual logo de fornecedor vai na proforma, a partir dos itens."""
        catalogo_fornecedores = self.config.fornecedores
        if (
            resolvido.fornecedores_misturados
            and catalogo_fornecedores.pedido_misto == "nenhuma"
        ):
            return None
        return catalogo_fornecedores.get(resolvido.fornecedor_dominante)

    def gerar(self, resolvido: PedidoResolvido, destino: Path | str) -> Path:
        wb = Workbook()
        wb.remove(wb.active)

        fornecedor = self.fornecedor_do_pedido(resolvido)
        controle = self._aba_controle(wb, resolvido)
        infos = [
            self._aba_proforma(
                wb, resolvido, proforma, controle, len(resolvido.proformas), fornecedor
            )
            for proforma in resolvido.proformas
        ]
        self._conciliar(controle, infos)

        pedido = resolvido.pedido
        wb.properties.title = f"Proforma {pedido.referencia}".strip()
        wb.properties.creator = self.empresa.razao_social or "augeo-compras"

        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destino)
        return destino


def nome_sugerido(resolvido: PedidoResolvido, quando: date | None = None) -> str:
    """Nome de arquivo previsível: data + referência/número da proforma."""
    pedido = resolvido.pedido
    quando = quando or pedido.data or date.today()
    numero = resolvido.proformas[0].numero if resolvido.proformas else ""
    partes = [quando.strftime("%Y-%m-%d"), pedido.referencia or numero or "proforma"]
    bruto = "_".join(p for p in partes if p)
    limpo = "".join(c if c.isalnum() or c in "-_. " else "-" for c in bruto).strip()
    return f"{limpo.replace(' ', '_')}.xlsx"


def gerar(config: Config, resolvido: PedidoResolvido, destino: Path | str) -> Path:
    """Atalho funcional para `GeradorPlanilha(config).gerar(...)`."""
    return GeradorPlanilha(config).gerar(resolvido, destino)
