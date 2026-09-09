"""Camada de serviço: junta configuração, catálogo, pedido e planilha.

É a porta de entrada única usada tanto pela linha de comando quanto pela
interface web, para que as duas se comportem exatamente da mesma forma.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from functools import lru_cache
from pathlib import Path
from typing import Any

from . import catalogo as mod_catalogo
from . import pedido as mod_pedido
from . import planilha as mod_planilha
from .catalogo import Catalogo, Item
from .config import Config, Empresa, Fornecedor, carregar as carregar_config
from .pedido import Pedido, PedidoResolvido
from .precos import PrecoCalculado, TabelaDePrecos


@dataclass
class ItemComPreco:
    """Item do catálogo já com o preço líquido calculado — para exibição."""

    item: Item
    preco: PrecoCalculado

    def para_dict(self) -> dict[str, Any]:
        return {
            **self.item.para_dict(),
            "desconto_percentual": self.preco.desconto_percentual,
            "preco_liquido": self.preco.preco_liquido,
            "fonte_desconto": self.preco.fonte_desconto,
        }


class Servico:
    """Ponto único de acesso ao sistema."""

    def __init__(self, config: Config | None = None):
        self.config = config or carregar_config()
        self._catalogo: Catalogo | None = None

    # -- catálogo ---------------------------------------------------------- #
    @property
    def catalogo(self) -> Catalogo:
        if self._catalogo is None:
            self._catalogo = mod_catalogo.carregar(self.config)
        return self._catalogo

    def recarregar_catalogo(self) -> Catalogo:
        self._catalogo = mod_catalogo.carregar(self.config)
        return self._catalogo

    @property
    def tabela(self) -> TabelaDePrecos:
        return TabelaDePrecos(self.config.comercial)

    def com_preco(self, item: Item) -> ItemComPreco:
        return ItemComPreco(item=item, preco=self.tabela.calcular(item))

    def buscar(
        self,
        termo: str = "",
        *,
        sistema: str | None = None,
        categoria: str | None = None,
        limite: int | None = 200,
    ) -> list[ItemComPreco]:
        encontrados = self.catalogo.buscar(
            termo, sistema=sistema, categoria=categoria, limite=limite
        )
        return [self.com_preco(item) for item in encontrados]

    def exportar_catalogo(self, destino: Path | str) -> Path:
        """Gera um Excel com todo o catálogo: preço bruto, desconto e líquido."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Catálogo"
        cabecalho = [
            "Sistema", "Categoria", "Part Number", "Type", "Description",
            f"Bruto {self.config.comercial.moeda}", "Desconto %",
            f"Líquido {self.config.comercial.moeda}", "Arquivo",
        ]
        ws.append(cabecalho)
        for coluna, _ in enumerate(cabecalho, start=1):
            celula = ws.cell(row=1, column=coluna)
            celula.font = Font(bold=True)
            celula.alignment = Alignment(horizontal="center")

        tabela = self.tabela
        for item in self.catalogo:
            calculo = tabela.calcular(item)
            ws.append([
                item.sistema, item.categoria, item.part_number, item.tipo, item.descricao,
                item.preco_bruto, calculo.desconto_percentual, calculo.preco_liquido, item.origem,
            ])

        for coluna, largura in zip("ABCDEFGHI", (22, 30, 20, 20, 70, 14, 11, 14, 24)):
            ws.column_dimensions[coluna].width = largura
        for linha in range(2, ws.max_row + 1):
            for coluna in ("F", "H"):
                ws[f"{coluna}{linha}"].number_format = "#,##0.00"
            ws[f"G{linha}"].number_format = "0.0"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:I{ws.max_row}"

        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destino)
        return destino

    # -- pedido ------------------------------------------------------------ #
    def resolver(self, pedido: Pedido) -> PedidoResolvido:
        return mod_pedido.resolver(pedido, self.catalogo, self.config.comercial)

    def resumo(self, resolvido: PedidoResolvido) -> dict[str, Any]:
        """Números e alertas que a interface mostra — iguais no servidor e no
        navegador, por virem daqui."""
        fornecedor = mod_planilha.GeradorPlanilha(self.config).fornecedor_do_pedido(resolvido)
        return {
            "itens": len(resolvido.linhas),
            "fornecedor": (
                {
                    "id": fornecedor.id,
                    "nome": fornecedor.nome,
                    "tem_logo": bool(fornecedor.caminho_logo),
                }
                if fornecedor
                else None
            ),
            "fornecedores_misturados": resolvido.fornecedores_misturados,
            "total_controle": resolvido.total_controle,
            "frete_total": resolvido.frete_total,
            "valor_faturado": resolvido.valor_faturado,
            "valor_servico": resolvido.valor_servico,
            "total_geral": resolvido.total_geral,
            "avisos": resolvido.avisos,
            "proformas": [
                {
                    "indice": p.indice,
                    "numero": p.numero,
                    "itens": len(p.linhas),
                    "subtotal": p.subtotal,
                    "frete": p.frete,
                    "total": p.total,
                }
                for p in resolvido.proformas
            ],
            "linhas": [
                {
                    "posicao": linha.posicao,
                    "part_number": linha.part_number,
                    "tipo": linha.tipo,
                    "descricao": linha.descricao,
                    "quantidade": linha.quantidade,
                    "preco_bruto": linha.preco_bruto,
                    "desconto_percentual": linha.desconto_percentual,
                    "preco_unitario": linha.preco_unitario,
                    "total": linha.total,
                    "quantidade_proforma": linha.quantidade_proforma,
                    "preco_proforma": linha.preco_proforma,
                    "total_proforma": linha.total_proforma,
                    "divergente": linha.divergente,
                    "proforma": linha.proforma,
                }
                for linha in resolvido.linhas
            ],
        }

    # -- identidade (quem compra e de quem) -------------------------------- #
    def _guardar_logo(self, conteudo_base64: str, nome: str) -> str:
        """Grava uma logo enviada pela tela e devolve o caminho absoluto.

        Absoluto de propósito: caminho relativo seria resolvido a partir da raiz
        do projeto, e a pasta de logos acompanha a `raiz` desta configuração.
        """
        import base64

        pasta = self.config.pasta_logos
        pasta.mkdir(parents=True, exist_ok=True)
        seguro = "".join(c if c.isalnum() or c in "-_." else "-" for c in nome) or "logo.png"
        destino = pasta / seguro
        destino.write_bytes(base64.b64decode(conteudo_base64))
        return str(destino)

    def _logo_em_dados(self, caminho: Path | None) -> str | None:
        """Devolve a logo como data URI, para a interface exibir."""
        import base64
        import mimetypes

        if not caminho or not caminho.exists():
            return None
        tipo = mimetypes.guess_type(caminho.name)[0] or "image/png"
        return f"data:{tipo};base64,{base64.b64encode(caminho.read_bytes()).decode()}"

    def aplicar_identidade(self, dados: dict[str, Any]) -> dict[str, Any]:
        """Troca o emitente em uso e o fornecedor, incluindo as logos enviadas.

        `remover` apaga emitentes pelo identificador; um emitente cujo
        identificador ainda não existe é criado.
        """
        for identificador in dados.get("remover") or []:
            self.config.emitentes = [
                e for e in self.config.emitentes if e.identificador != identificador
            ]

        entrada = dados.get("emitente") or {}
        if entrada:
            atual = self.config.emitente(entrada.get("identificador", "")) or Empresa(
                identificador=entrada.get("identificador", "") or "emitente"
            )
            campos = {**atual.para_dict(), **{k: v for k, v in entrada.items() if k in atual.para_dict()}}
            if entrada.get("logo_conteudo"):
                campos["logo"] = self._guardar_logo(
                    entrada["logo_conteudo"], entrada.get("logo_nome") or "emitente.png"
                )
            empresa = Empresa.de_dict(campos)
            empresa.largura_logo = atual.largura_logo
            self.config.empresa = empresa
            outros = [e for e in self.config.emitentes if e.identificador != empresa.identificador]
            self.config.emitentes = [*outros, empresa] if empresa.identificador else outros

        if not self.config.emitentes:
            self.config.emitentes = [self.config.empresa]

        entrada = dados.get("fornecedor") or {}
        if entrada:
            catalogo = self.config.fornecedores
            identificador = entrada.get("id") or catalogo.padrao or "fornecedor"
            atual = catalogo.itens.get(identificador) or Fornecedor(id=identificador)
            if entrada.get("nome"):
                atual.nome = entrada["nome"]
            if entrada.get("logo_conteudo"):
                atual.logo = self._guardar_logo(
                    entrada["logo_conteudo"], entrada.get("logo_nome") or "fornecedor.png"
                )
            atual.id = identificador
            catalogo.itens[identificador] = atual
            catalogo.padrao = identificador

        return self.identidade()

    def identidade(self) -> dict[str, Any]:
        """Quem está comprando, de quem, e as logos — para a interface mostrar."""
        empresa = self.config.empresa
        fornecedor = self.config.fornecedores.get(None)
        return {
            "emitente": empresa.para_dict(),
            "emitentes": [
                {
                    "identificador": e.identificador,
                    "razao_social": e.razao_social,
                    "cnpj": e.cnpj,
                }
                for e in self.config.emitentes
            ],
            "logo_emitente": self._logo_em_dados(empresa.caminho_logo),
            "fornecedor": (
                {"id": fornecedor.id, "nome": fornecedor.nome} if fornecedor else None
            ),
            "logo_fornecedor": self._logo_em_dados(fornecedor.caminho_logo) if fornecedor else None,
            # O CNPJ é opcional: não vem de lugar nenhum automaticamente, e
            # nem toda proforma precisa dele.
            "completo": bool(empresa.razao_social),
        }

    def descricao_config(self) -> dict[str, Any]:
        """O que a interface precisa saber da configuração ao iniciar."""
        comercial, empresa = self.config.comercial, self.config.empresa
        return {
            "moeda": comercial.moeda,
            "simbolo_moeda": comercial.simbolo_moeda,
            "desconto_percentual": comercial.desconto_percentual,
            "condicoes_padrao": comercial.condicoes_padrao.para_dict(),
            "empresa": {"razao_social": empresa.razao_social, "cidade": empresa.cidade},
            "sistemas": self.catalogo.sistemas,
            "total_itens": len(self.catalogo),
            "avisos_catalogo": self.catalogo.avisos,
            "fornecedores": [
                {"id": f.id, "nome": f.nome, "tem_logo": bool(f.caminho_logo)}
                for f in self.config.fornecedores.itens.values()
            ],
        }

    def gerar(
        self,
        pedido: Pedido,
        destino: Path | str | None = None,
    ) -> tuple[Path, PedidoResolvido]:
        """Resolve o pedido e grava a planilha. Devolve (caminho, resolvido)."""
        resolvido = self.resolver(pedido)
        if destino is None:
            destino = self.config.pasta_saida / mod_planilha.nome_sugerido(resolvido)
        caminho = mod_planilha.gerar(self.config, resolvido, destino)
        return caminho, resolvido

    def pedido_de_planilha(self, caminho: Path | str, **cabecalho: Any) -> Pedido:
        return Pedido.de_planilha(caminho, **cabecalho)

    # -- pedidos salvos ---------------------------------------------------- #
    def listar_pedidos(self) -> list[Path]:
        pasta = self.config.pasta_pedidos
        if not pasta.is_dir():
            return []
        return sorted(pasta.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True)

    def salvar_pedido(self, pedido: Pedido, nome: str | None = None) -> Path:
        base = nome or pedido.referencia or date.today().isoformat()
        seguro = "".join(c if c.isalnum() or c in "-_." else "-" for c in base).strip("-")
        return pedido.salvar(self.config.pasta_pedidos / f"{seguro or 'pedido'}.json")


@lru_cache(maxsize=1)
def servico_padrao() -> Servico:
    """Instância compartilhada (o catálogo é carregado uma única vez)."""
    return Servico()
