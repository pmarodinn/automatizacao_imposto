"""Geração do arquivo Excel: estrutura, fórmulas e recálculo dos totais."""

from datetime import date

import openpyxl
import pytest

from augeo_compras.pedido import LinhaPedido, Pedido, Proforma
from augeo_compras.planilha import GeradorPlanilha, nome_sugerido
from augeo_compras.servico import Servico


def _linha_do_rotulo(ws, rotulo: str) -> int:
    """Número da linha onde um rótulo aparece — evita depender de coordenadas."""
    for linha in ws.iter_rows():
        for celula in linha:
            if celula.value == rotulo:
                return celula.row
    raise AssertionError(f"rótulo {rotulo!r} não encontrado")


def _valor_ao_lado(ws, rotulo: str):
    """Encontra o rótulo de um total e devolve o valor/fórmula da célula ao lado."""
    for linha in ws.iter_rows():
        for celula in linha:
            if celula.value == rotulo:
                return ws.cell(row=celula.row, column=celula.column + 1).value
    raise AssertionError(f"rótulo {rotulo!r} não encontrado")


@pytest.fixture
def servico(config, catalogo_real):
    servico = Servico(config)
    servico._catalogo = catalogo_real
    return servico


@pytest.fixture
def pedido_exemplo():
    """Reproduz o pedido do arquivo modelo, incluindo a divergência de 1 peça."""
    return Pedido(
        referencia="TRT",
        cliente="Securiton AG",
        data=date(2026, 8, 29),
        proformas=[Proforma(numero="2608261-1", data=date(2026, 8, 29))],
        linhas=[
            LinhaPedido("11-2000001-01-04", 2),
            LinhaPedido("11-2000017-01-02", 2),
            LinhaPedido("11-2000018-01-02", 1),
            LinhaPedido("11-2000010-01-02", 4),
            LinhaPedido("11-2300030-01-02", 5, quantidade_proforma=6),
            LinhaPedido("50-0500845-01-01", 17),
            LinhaPedido("50-0500844-01-01", 1),
        ],
    )


def test_abas_e_cabecalhos(servico, pedido_exemplo, tmp_path):
    caminho, _ = servico.gerar(pedido_exemplo, tmp_path / "saida.xlsx")
    wb = openpyxl.load_workbook(caminho)

    assert wb.sheetnames == ["NET PRICE", "PROFORMA"]
    controle = wb["NET PRICE"]
    assert controle["A1"].value == "CONTROL"
    assert [controle[c].value for c in ("A3", "C3", "D3", "G3", "I3", "J3", "K3")] == [
        "Pos", "Quantity", "Part Number", "Description", "Code", "Preço EUR", "Saldo EUR"
    ]


def test_formulas_ligam_as_duas_abas(servico, pedido_exemplo, tmp_path):
    caminho, _ = servico.gerar(pedido_exemplo, tmp_path / "saida.xlsx")
    wb = openpyxl.load_workbook(caminho)
    controle, proforma = wb["NET PRICE"], wb["PROFORMA"]

    assert controle["K4"].value == "=J4*C4"
    assert controle["K11"].value == "=SUM(K4:K10)"

    # A proforma puxa da aba de controle...
    primeira = _linha_do_rotulo(proforma, "Pos") + 1
    assert proforma.cell(row=primeira, column=3).value == "='NET PRICE'!C4"
    assert proforma.cell(row=primeira, column=10).value == "='NET PRICE'!J4"
    # ...menos onde declara quantidade diferente (5 comprados, 6 na proforma).
    assert proforma.cell(row=primeira + 4, column=3).value == 6

    # E a conciliação volta da proforma para o controle.
    total_proforma = _linha_do_rotulo(proforma, "TOTAL")
    assert _valor_ao_lado(controle, "Invoice Value") == f"='PROFORMA'!J{total_proforma}"
    assert _valor_ao_lado(controle, "Service Value") == "=K11-J14+J12"


def test_posicoes_e_precos_liquidos(servico, pedido_exemplo, tmp_path):
    caminho, resolvido = servico.gerar(pedido_exemplo, tmp_path / "saida.xlsx")
    wb = openpyxl.load_workbook(caminho)
    controle = wb["NET PRICE"]

    assert [controle[f"A{linha}"].value for linha in range(4, 11)] == [10, 20, 30, 40, 50, 60, 70]
    assert controle["D4"].value == "11-2000001-01-04"
    assert controle["J4"].value == 664.89  # 1198,00 com 44,5% de desconto
    assert resolvido.total_controle == 5216.44
    assert resolvido.valor_servico == -67.15


def test_coluna_auxiliar_fica_oculta(servico, pedido_exemplo, tmp_path):
    caminho, _ = servico.gerar(pedido_exemplo, tmp_path / "saida.xlsx")
    wb = openpyxl.load_workbook(caminho)
    controle = wb["NET PRICE"]
    assert controle["L4"].value == "=J4*0.4"
    assert controle.column_dimensions["L"].hidden is True


def test_timbre_e_condicoes_na_proforma(servico, pedido_exemplo, tmp_path):
    caminho, _ = servico.gerar(pedido_exemplo, tmp_path / "saida.xlsx")
    wb = openpyxl.load_workbook(caminho)
    proforma = wb["PROFORMA"]

    textos = [c.value for linha in proforma.iter_rows(max_row=11) for c in linha if c.value]
    assert "PARANÁ EM REDE SISTEMA LTDA" in textos
    assert "PROFORMA:" in textos
    assert "2608261-1" in textos
    assert any("29 de agosto de 2026" in str(t) for t in textos)

    rodape = [c.value for linha in proforma.iter_rows(min_row=23) for c in linha if c.value]
    assert "Condições gerais" in rodape
    assert "Pagamento em 60 dias" in rodape


def test_uma_aba_por_proforma(servico, tmp_path):
    pedido = Pedido(
        referencia="SPLIT",
        proformas=[Proforma(numero="A-1"), Proforma(numero="A-2", frete=80.0)],
        linhas=[
            LinhaPedido("11-2000001-01-04", 1, proforma=1),
            LinhaPedido("11-2300030-01-02", 2, proforma=2),
        ],
    )
    caminho, resolvido = servico.gerar(pedido, tmp_path / "split.xlsx")
    wb = openpyxl.load_workbook(caminho)

    assert wb.sheetnames == ["NET PRICE", "PROFORMA 1", "PROFORMA 2"]
    assert resolvido.frete_total == 80.0

    faturado = _valor_ao_lado(wb["NET PRICE"], "Invoice Value")
    assert "'PROFORMA 1'!" in faturado
    assert "'PROFORMA 2'!" in faturado


def test_pedido_sem_itens_nao_quebra(servico, tmp_path):
    caminho, _ = servico.gerar(Pedido(referencia="VAZIO"), tmp_path / "vazio.xlsx")
    wb = openpyxl.load_workbook(caminho)
    assert wb["NET PRICE"]["A1"].value == "CONTROL"


def test_linhas_reservadas_geram_linhas_em_branco(config, catalogo_real, tmp_path):
    import copy

    cfg = copy.deepcopy(config)
    cfg.comercial.linhas_reservadas = 10
    servico = Servico(cfg)
    servico._catalogo = catalogo_real

    pedido = Pedido(linhas=[LinhaPedido("11-2300030-01-02", 1)])
    caminho, _ = servico.gerar(pedido, tmp_path / "reservadas.xlsx")
    controle = openpyxl.load_workbook(caminho)["NET PRICE"]

    assert controle["A13"].value == 100  # 10 linhas: posições 10..100
    assert controle["D13"].value is None
    assert controle["K13"].value == "=J13*C13"


def test_nome_sugerido(servico, pedido_exemplo):
    resolvido = servico.resolver(pedido_exemplo)
    assert nome_sugerido(resolvido) == "2026-08-29_TRT.xlsx"


def test_layout_e_dirigido_pela_configuracao(config):
    """Trocar o layout no YAML muda as colunas da planilha, sem tocar em código."""
    import copy

    cfg = copy.deepcopy(config)
    cfg.layout.colunas = [c for c in cfg.layout.colunas if c.chave != "codigo"]
    gerador = GeradorPlanilha(cfg)

    assert "codigo" not in gerador.posicoes
    assert gerador.total_colunas == 10
