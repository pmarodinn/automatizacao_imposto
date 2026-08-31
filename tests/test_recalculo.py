"""Confere que as fórmulas gravadas realmente calculam os valores esperados.

Usa o LibreOffice para recalcular a planilha e compara com os totais que o
Python já havia computado. É o teste que garante que a planilha entregue ao
usuário fecha as contas sozinha. Pulado quando o LibreOffice não está instalado.
"""

from __future__ import annotations

import shutil
import subprocess
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from augeo_compras.pedido import LinhaPedido, Pedido, Proforma
from augeo_compras.servico import Servico

soffice = shutil.which("soffice") or shutil.which("libreoffice")
pytestmark = pytest.mark.skipif(soffice is None, reason="LibreOffice não instalado")


def recalcular(caminho: Path, saida: Path) -> openpyxl.Workbook:
    """Reabre a planilha no LibreOffice para que ele avalie as fórmulas."""
    subprocess.run(
        [soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(saida), str(caminho)],
        check=True,
        capture_output=True,
        timeout=180,
    )
    return openpyxl.load_workbook(saida / caminho.name, data_only=True)


def valor(ws, rotulo: str) -> float:
    for linha in ws.iter_rows():
        for celula in linha:
            if celula.value == rotulo:
                for coluna in range(celula.column + 1, ws.max_column + 1):
                    lido = ws.cell(row=celula.row, column=coluna).value
                    if isinstance(lido, (int, float)):
                        return float(lido)
    raise AssertionError(f"valor de {rotulo!r} não encontrado")


@pytest.mark.slow
def test_totais_recalculados_batem_com_o_python(config, catalogo_real, tmp_path):
    servico = Servico(config)
    servico._catalogo = catalogo_real

    pedido = Pedido(
        referencia="RECALCULO",
        data=date(2026, 8, 29),
        proformas=[Proforma(numero="A-1", frete=150.0), Proforma(numero="A-2")],
        linhas=[
            LinhaPedido("11-2000001-01-04", 2, proforma=1),
            LinhaPedido("11-2300030-01-02", 5, quantidade_proforma=6, proforma=1),
            LinhaPedido("20-1151010-01-01", 3, proforma=2),
        ],
    )
    caminho, resolvido = servico.gerar(pedido, tmp_path / "recalculo.xlsx")
    wb = recalcular(caminho, tmp_path / "recalculado")

    controle = wb["NET PRICE"]
    assert valor(controle, "TOTAL") == pytest.approx(resolvido.total_controle, abs=0.01)
    assert valor(controle, "FREIGHT") == pytest.approx(resolvido.frete_total, abs=0.01)
    assert valor(controle, "Invoice Value") == pytest.approx(resolvido.valor_faturado, abs=0.01)
    assert valor(controle, "Service Value") == pytest.approx(resolvido.valor_servico, abs=0.01)
    assert valor(controle, "Total") == pytest.approx(resolvido.total_geral, abs=0.01)

    for indice, proforma in enumerate(resolvido.proformas, start=1):
        ws = wb[f"PROFORMA {indice}"]
        assert valor(ws, "SUBTOTAL") == pytest.approx(proforma.subtotal, abs=0.01)
        assert valor(ws, "TOTAL") == pytest.approx(proforma.total, abs=0.01)
